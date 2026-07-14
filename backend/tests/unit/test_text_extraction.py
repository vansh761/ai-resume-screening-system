"""
Unit tests for text extraction.

We generate real, minimal PDF and DOCX files in-memory rather than
mocking `pdfplumber`/`python-docx` — mocking the parsing library would
only prove our code calls the mock correctly, not that extraction
actually works against a real file. `fpdf2` (test-fixture-only, see
requirements.txt) and `python-docx` (already a real dependency, used
here to construct rather than read) both let us build valid files
without needing a repo of sample resumes checked into version control.
"""

import io

import pytest
from docx import Document
from fpdf import FPDF

from app.ai.text_extraction import (
    TextExtractionError,
    extract_text,
    extract_text_from_docx,
    extract_text_from_pdf,
)
from app.models.resume import FileType


def _make_pdf_bytes(text: str) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text=text)
    return bytes(pdf.output())


def _make_docx_bytes(paragraphs: list[str]) -> bytes:
    document = Document()
    for para in paragraphs:
        document.add_paragraph(para)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.mark.unit
def test_extract_text_from_pdf_returns_expected_content() -> None:
    pdf_bytes = _make_pdf_bytes("Jane Doe Software Engineer")
    result = extract_text_from_pdf(pdf_bytes)
    assert "Jane Doe" in result


@pytest.mark.unit
def test_extract_text_from_docx_returns_all_paragraphs() -> None:
    docx_bytes = _make_docx_bytes(["John Smith", "Backend Engineer", "5 years experience"])
    result = extract_text_from_docx(docx_bytes)
    assert "John Smith" in result
    assert "Backend Engineer" in result
    assert "5 years experience" in result


@pytest.mark.unit
def test_extract_text_from_docx_skips_empty_paragraphs() -> None:
    docx_bytes = _make_docx_bytes(["Real content", "", "   ", "More content"])
    result = extract_text_from_docx(docx_bytes)
    lines = [line for line in result.split("\n") if line]
    assert len(lines) == 2


@pytest.mark.unit
def test_extract_text_from_pdf_raises_on_garbage_input() -> None:
    with pytest.raises(TextExtractionError):
        extract_text_from_pdf(b"this is not a pdf file at all")


@pytest.mark.unit
def test_extract_text_from_docx_raises_on_garbage_input() -> None:
    with pytest.raises(TextExtractionError):
        extract_text_from_docx(b"this is not a docx file at all")


@pytest.mark.unit
def test_extract_text_dispatches_by_file_type() -> None:
    docx_bytes = _make_docx_bytes(["Dispatch test content"])
    result = extract_text(docx_bytes, FileType.DOCX)
    assert "Dispatch test content" in result
